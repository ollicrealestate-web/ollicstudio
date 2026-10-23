"""Construction schedules: read the programme the contractor already has, and match each task to
a part of the architect's model, so the construction animation follows the real order and dates.

Read here (no licence needed, nothing sent to any AI):
- Microsoft Project XML (File > Save As > XML), tasks with Name, Start, Finish, Summary, OutlineLevel
- Primavera P6 XER export, TASK table (actual dates, then planned, then early), PROJWBS names
- CSV, and Excel .xlsx (for example SYNCHRO, Asta or Excel exports): columns found by their headings
  (task/activity/name, start, finish/end, optional WBS/phase/location)

Matching is by keywords in the task name and its parent WBS (e.g. "Level 2 > Slab"): the building
part (foundation, columns, slabs, walls, windows, roof, MEP...) and the floor. The customer can
correct any match in the app before the animation is made.
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, timedelta
from xml.etree import ElementTree as ET

MAX_TASKS = 400
GROUPS = ("site", "foundation", "structure", "walls", "openings", "roof", "mep", "finishes", "furnish", "none")
KINDS = ("columns", "beams", "slabs", "stairs")
GROUP_NAMES = {"site": "Site works", "foundation": "Foundation", "structure": "Structure", "walls": "Walls and facade", "openings": "Doors and windows",
               "roof": "Roof", "mep": "MEP services", "finishes": "Finishes", "furnish": "Furniture and fittings", "none": "Not shown (no building part)"}


class ScheduleUnreadable(Exception):
    """`reason` is safe to show the customer."""

    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


@dataclass
class Task:
    name: str
    start: date
    finish: date
    context: str = ""            # parent WBS / summary names, used for matching
    group: str = "none"
    kind: str | None = None
    floors: list[str] = field(default_factory=list)
    milestone: bool = False

    def public(self, i: int) -> dict:
        d = asdict(self)
        return {"id": i, **{k: d[k] for k in ("name", "group", "kind", "floors", "milestone", "context")},
                "start": self.start.isoformat(), "finish": self.finish.isoformat()}


# ---------------------------------------------------------------- dates
MONTHS = {m: i + 1 for i, m in enumerate(("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"))}


def _year(y: int) -> int:
    return y + 2000 if y < 100 else y


def parse_date(v, day_first: bool = True) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):          # an Excel serial number
        return (datetime(1899, 12, 30) + timedelta(days=float(v))).date() if 20000 < v < 80000 else None
    s = str(v).strip()
    s = re.sub(r"^(mon|tue|wed|thu|fri|sat|sun)[a-z]*\.?,?\s+", "", s, flags=re.I)   # "Mon 01/03/27" from MS Project
    s = re.sub(r"\s*[AP]M?$|\*$", "", s, flags=re.I).strip()
    m = re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        return _safe(int(m[1]), int(m[2]), int(m[3]))
    m = re.match(r"(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})", s)
    if m:
        a, b, y = int(m[1]), int(m[2]), _year(int(m[3]))
        d, mo = (a, b) if day_first else (b, a)
        return _safe(y, mo, d) or _safe(y, d, mo)
    m = re.match(r"(\d{1,2})[\s\-]([A-Za-z]{3})[A-Za-z]*[\s\-,]+(\d{2,4})", s)          # 1-Mar-27, 01 March 2027
    if m and m[2].lower() in MONTHS:
        return _safe(_year(int(m[3])), MONTHS[m[2].lower()], int(m[1]))
    m = re.match(r"([A-Za-z]{3})[A-Za-z]*\.?\s+(\d{1,2}),?\s+(\d{2,4})", s)            # March 1, 2027
    if m and m[1].lower() in MONTHS:
        return _safe(_year(int(m[3])), MONTHS[m[1].lower()], int(m[2]))
    return None


def _safe(y, m, d) -> date | None:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def _day_first(values: list[str]) -> bool:
    """dd/mm unless the dates can only be mm/dd (a second number above 12)."""
    for v in values:
        m = re.search(r"\b(\d{1,2})[/.\-](\d{1,2})[/.\-]\d{2,4}", str(v))
        if m and int(m[1]) > 12:
            return True
        if m and int(m[2]) > 12:
            return False
    return True


# ---------------------------------------------------------------- formats
def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _msproject(data: bytes) -> list[Task]:
    try:
        root = ET.fromstring(data)
    except ET.ParseError as exc:
        raise ScheduleUnreadable("This XML file is damaged or incomplete. Save it from Microsoft Project again (File > Save As > XML).") from exc
    tasks, parents = [], {}
    for t in root.iter():
        if _local(t.tag) != "Task":
            continue
        f = {_local(c.tag): (c.text or "").strip() for c in t}
        name, level = f.get("Name", ""), int(f.get("OutlineLevel") or 1) if (f.get("OutlineLevel") or "1").isdigit() else 1
        if not name:
            continue
        parents[level] = name
        for k in [k for k in parents if k > level]:
            del parents[k]
        if f.get("Summary") == "1" or f.get("IsNull") == "1":
            continue
        start, finish = parse_date(f.get("ActualStart")) or parse_date(f.get("Start")), parse_date(f.get("ActualFinish")) or parse_date(f.get("Finish"))
        if start and finish:
            ctx = " > ".join(parents[k] for k in sorted(parents) if k < level)
            tasks.append(Task(name=name, start=start, finish=max(start, finish), context=ctx, milestone=f.get("Milestone") == "1"))
    return tasks


def _xer(text: str) -> list[Task]:
    tables, cur, fields = {}, None, []
    for line in text.splitlines():
        parts = line.split("\t")
        if parts[0] == "%T":
            cur = parts[1].strip() if len(parts) > 1 else None
            tables.setdefault(cur, [])
        elif parts[0] == "%F":
            fields = [p.strip() for p in parts[1:]]
        elif parts[0] == "%R" and cur:
            tables[cur].append(dict(zip(fields, parts[1:])))
    wbs = {w.get("wbs_id"): w for w in tables.get("PROJWBS", [])}

    def path(wid):
        names, seen = [], set()
        while wid in wbs and wid not in seen:
            seen.add(wid)
            w = wbs[wid]
            if w.get("proj_node_flag") != "Y":
                names.append(w.get("wbs_name", ""))
            wid = w.get("parent_wbs_id")
        return " > ".join(reversed([n for n in names if n]))

    out = []
    for r in tables.get("TASK", []):
        name = (r.get("task_name") or "").strip()
        start = next((d for d in (parse_date(r.get(k), day_first=False) for k in ("act_start_date", "target_start_date", "early_start_date", "restart_date")) if d), None)
        finish = next((d for d in (parse_date(r.get(k), day_first=False) for k in ("act_end_date", "target_end_date", "early_end_date", "reend_date")) if d), None)
        if name and start and finish:
            out.append(Task(name=name, start=start, finish=max(start, finish), context=path(r.get("wbs_id")),
                            milestone=(r.get("task_type") or "") in ("TT_Mile", "TT_FinMile")))
    return out


HEAD = {"name": ("task name", "activity name", "task", "activity", "name", "description", "title"),
        "start": ("actual start", "start date", "start", "planned start", "baseline start", "early start", "begin"),
        "finish": ("actual finish", "finish date", "end date", "finish", "end", "planned finish", "planned end", "baseline finish", "early finish", "due"),
        "context": ("wbs name", "wbs", "phase", "location", "zone", "area", "level", "summary", "parent")}


def _rows(rows: list[list]) -> list[Task]:
    rows = [[c if not isinstance(c, str) else c.strip() for c in r] for r in rows if any(c not in (None, "") for c in r)]
    for hi, header in enumerate(rows[:15]):          # the heading row may sit under a title
        low = [str(c or "").strip().lower() for c in header]
        col = {}
        for key, names in HEAD.items():
            for n in names:
                if n in low and low.index(n) not in col.values():
                    col[key] = low.index(n); break
        if {"name", "start", "finish"} <= set(col):
            break
    else:
        raise ScheduleUnreadable("No task, start and finish columns were found. The first row should have headings such as Task Name, Start and Finish.")
    body = rows[hi + 1:]
    day_first = _day_first([r[col["start"]] for r in body if col["start"] < len(r)])
    out = []
    for r in body:
        get = lambda k: r[col[k]] if k in col and col[k] < len(r) else None
        name = " ".join(str(get("name") or "").split())
        start, finish = parse_date(get("start"), day_first), parse_date(get("finish"), day_first)
        if name and start and finish:
            out.append(Task(name=name, start=start, finish=max(start, finish), context=" ".join(str(get("context") or "").split())))
    return out


def _xlsx(data: bytes) -> list[Task]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ScheduleUnreadable("This Excel file could not be opened. Save it again as .xlsx, or export the schedule as CSV.") from exc
    last = None
    for ws in wb.worksheets:
        try:
            tasks = _rows([list(r) for r in ws.iter_rows(max_row=5000, values_only=True)])
        except ScheduleUnreadable as exc:
            last = exc
            continue
        if tasks:
            return tasks
    raise last or ScheduleUnreadable("No tasks with dates were found in this spreadsheet.")


def read_schedule(data: bytes, filename: str) -> tuple[str, list[Task]]:
    """Returns (format name, tasks with dates). Summary rows are left out; at most MAX_TASKS are kept."""
    if not data:
        raise ScheduleUnreadable("The file is empty.")
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    head = data[:4000].lstrip(b"\xef\xbb\xbf \r\n\t")
    if data.startswith(b"PK\x03\x04"):
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile as exc:
            raise ScheduleUnreadable("This file is damaged. Save or export it again.") from exc
        if not any(n.startswith("xl/") for n in names):
            raise ScheduleUnreadable("This is a ZIP file, not a schedule. Upload the .xml, .xer, .csv or .xlsx file itself.")
        fmt, tasks = "Excel", _xlsx(data)
    elif head.startswith(b"ERMHDR") or ext == "xer":
        fmt, tasks = "Primavera P6", _xer(data.decode("utf-8", errors="replace") if _utf8(data) else data.decode("cp1252", errors="replace"))
    elif head.startswith(b"<") and b"<Project" in data[:5000]:
        fmt, tasks = "Microsoft Project", _msproject(data)
    elif ext == "mpp" or data.startswith(b"\xd0\xcf\x11\xe0"):
        raise ScheduleUnreadable("Microsoft Project .mpp files cannot be read directly. In Project use File > Save As > XML, and upload the .xml file.")
    else:
        try:
            text = data.decode("utf-8-sig") if _utf8(data) else data.decode("cp1252")
        except UnicodeDecodeError as exc:
            raise ScheduleUnreadable("This file type cannot be read. Upload a Microsoft Project .xml, Primavera .xer, .csv or .xlsx schedule.") from exc
        try:
            dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        fmt, tasks = "CSV", _rows(list(csv.reader(io.StringIO(text), dialect)))
    if not tasks:
        raise ScheduleUnreadable(f"No tasks with start and finish dates were found in this {fmt} file.")
    tasks.sort(key=lambda t: (t.start, t.finish))
    return fmt, tasks[:MAX_TASKS]


def _utf8(data: bytes) -> bool:
    try:
        data[:20000].decode("utf-8")
        return True
    except UnicodeDecodeError as exc:
        return exc.start > 19990     # cut in the middle of a character at the end of the sample


# ---------------------------------------------------------------- matching
RULES = [  # (group, kind, pattern) first match wins; checked against the task name, then its WBS
    ("none", None, r"\b(handover|hand over|snag\w*|commission\w*|design|approval|permit|procure\w*|tender|submittal|mobili[sz]\w*|demobili[sz]\w*|award|contract|inspection|testing)\b"),
    ("roof", None, r"\b(roof\w*|parapet|penthouse roof|waterproofing to roof)\b"),
    ("foundation", None, r"\b(foundation\w*|footing\w*|pile\w*|piling|pile caps?|raft|substructure|blinding|ground beams?|oversite|dpm|ground (floor )?slab|base slab|strip footings?)\b"),
    ("site", None, r"\b(site (clearance|prep\w*|works|establishment)|setting out|set out|clear\w* (the )?site|excavat\w*|earth ?works?|survey\w*|demolition|hoarding|fenc\w*|landscap\w*|external works|paving|road\w*|car ?park\w*)\b"),
    ("openings", None, r"\b(windows?|doors?|glazing|curtain ?wall\w*|shop ?fronts?|louvres?|louvers?|skylights?)\b"),
    ("mep", None, r"\b(mep|m ?& ?e|mechanical|electrical|plumbing|hvac|ducts?|ductwork|pip(e|es|ing)|drainage|sprinklers?|fire alarm|cabl\w*|wiring|conduits?|lifts?|elevators?|escalators?|solar|generator|sewer\w*|services)\b"),
    ("furnish", None, r"\b(furniture|furnish\w*|ff ?& ?e|kitchen (units|fittings|cabinets)|sanitary\w*|wardrobes?|fixtures and fittings)\b"),
    ("finishes", None, r"\b(plaster\w*|screed\w*|til(e|es|ing)|paint\w*|ceilings?|finish\w*|floor(ing)? finish\w*|decorat\w*|fit[ -]?out|joinery|skirting|render(ing)? to)\b"),
    ("walls", None, r"\b(walls?|walling|block ?work|brick ?work|masonry|partitions?|cladding|fa[cç]ade|drywall|infill)\b"),
    ("structure", "stairs", r"\b(stair\w*|ramps?)\b"),
    ("structure", "columns", r"\b(columns?|pillars?|posts?)\b"),
    ("structure", "beams", r"\b(beams?|girders?|lintels?)\b"),
    ("structure", "slabs", r"\b(slabs?|deck\w*|suspended floor)\b"),
    ("structure", None, r"\b(frame|framing|structur\w*|superstructure|concrete|rc|reinforced|steel ?work|formwork|rebar|reinforcement|cast\w*|pour\w*|core)\b"),
]
ORD = {"ground": 0, "first": 1, "second": 2, "third": 3, "fourth": 4, "fifth": 5, "sixth": 6, "seventh": 7, "eighth": 8, "ninth": 9, "tenth": 10}


def level_of(text: str) -> list[int]:
    """Floor numbers named in the text: 'Level 2', 'L2', '2nd floor', 'first floor', 'GF', 'basement', 'levels 2-4'."""
    s = text.lower()
    m = re.search(r"\b(?:levels?|lvls?|floors?|storeys?|stor(?:ies|y))\s*(-?\d{1,3})\s*(?:-|–|to)\s*(?:l|level\s*)?(-?\d{1,3})\b", s) or \
        re.search(r"\bl(\d{1,3})\s*(?:-|–|to)\s*l(\d{1,3})\b", s)
    if m:
        a, b = int(m[1]), int(m[2])
        return list(range(min(a, b), max(a, b) + 1))[:60]
    m = re.search(r"\b(?:level|lvl|floor|storey|story|flr)\s*[-_ ]?(-?\d{1,3})\b", s) or re.search(r"\bl(-?\d{1,3})\b", s) or \
        re.search(r"\b(\d{1,3})(?:st|nd|rd|th)\s+(?:floor|level|storey)\b", s)
    if m:
        return [int(m[1])]
    m = re.search(r"\b(ground|first|second|third|fourth|fifth|sixth|seventh|eighth|ninth|tenth)\s+(?:floor|level|storey)\b", s)
    if m:
        return [ORD[m[1]]]
    if re.search(r"\b(gf|ground)\b", s):
        return [0]
    m = re.search(r"\b(?:b(\d)|basement\s*(\d)?)\b", s)
    if m:
        return [-int(m[1] or m[2] or 1)]
    return []


def storey_levels(storeys: list[str]) -> dict[str, int]:
    """A floor number for each storey of the model, from its name, else from its order (lowest first)."""
    named = {s: level_of(s) for s in storeys}
    if sum(1 for v in named.values() if v) >= max(1, len(storeys) // 2):
        out, last = {}, None
        for s in storeys:
            if named[s]:
                last = named[s][0]
            else:
                last = (last + 1) if last is not None else 0
            out[s] = last
        return out
    return {s: i for i, s in enumerate(storeys)}


def match_task(task: Task, storeys: list[str]) -> Task:
    name = task.name.lower()
    whole = f"{task.context} {task.name}".lower()
    hit = next(((g, k) for g, k, p in RULES if re.search(p, name)), None) or next(((g, k) for g, k, p in RULES if re.search(p, whole)), None)
    task.group, task.kind = hit if hit else ("none", None)
    if task.group == "none" and task.milestone:
        task.kind = None
    floors: list[str] = []
    if task.group not in ("site", "foundation", "roof", "none"):
        by_name = [s for s in storeys if s and re.search(rf"(?<!\w){re.escape(s.lower())}(?!\w)", whole)]
        if by_name:
            floors = by_name
        else:
            wanted = level_of(task.name) or level_of(task.context)
            if wanted:
                lv = storey_levels(storeys)
                floors = [s for s in storeys if lv.get(s) in wanted]
    task.floors = floors
    return task


def match_all(tasks: list[Task], storeys: list[str]) -> list[Task]:
    storeys = [" ".join(str(s).split())[:120] for s in storeys if str(s).strip()][:200]
    return [match_task(t, storeys) for t in tasks]
